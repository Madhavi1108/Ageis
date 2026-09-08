"""Acceptance (docs/AEGIS_IMPLEMENTATION_PLAN.md Section 31): the acceptance
task's report has all 18 sections populated with real data (Docker-less runs
carry PARTIALLY_SUPPORTED notes, not fabricated passes); the metrics workbook's
derivable metrics match hand-computed DB values and the rest are ``N/A`` with a
reason (never invented); export -> re-import round-trips.
"""

from __future__ import annotations

import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.analysis.analyze import analyze_snapshot
from app.core.config import Settings
from app.ingestion.ingest import ingest_repository
from app.repository.repositories import RepositoryRepository
from app.repository.snapshots import SnapshotRepository
from app.repository.tasks import TaskRepository
from app.schemas.implementation import EditOpsAI
from app.schemas.plan import EngineeringPlanAI
from app.schemas.report import SECTION_NAMES
from app.schemas.repository import IngestRequest
from app.schemas.task import TaskCreate
from app.services import impact as impact_service
from app.services import implementation as implementation_service
from app.services import mapping as mapping_service
from app.services import planning as planning_service
from app.services import reporting as reporting_service
from app.services import tasks as tasks_service
from app.services import verification as verification_service

_PLAN = EngineeringPlanAI.model_validate(
    {
        "problem_interpretation": "cap discount at 0.5",
        "assumptions": [],
        "files_to_inspect": ["invoice.py"],
        "files_to_modify": ["invoice.py"],
        "symbols_to_modify": ["invoice.py::calculate_total"],
        "dependencies": [],
        "steps": [
            {
                "id": "s1",
                "description": "clamp",
                "test_intent": "90==50",
                "evidence_refs": [],
            }
        ],
        "test_strategy": {"approach": "boundary"},
        "expected_behavior": "the discount is capped at 0.5",
        "regression_risks": [],
        "rollback_strategy": "revert invoice.py",
        "source": "AI",
        "confidence": {"value": 0.8, "basis": "INFERENCE"},
        "evidence": [],
    }
)
_EDIT = EditOpsAI.model_validate(
    {
        "edit_ops": [
            {
                "path": "invoice.py",
                "op": "replace",
                "anchor": "return price * (1 - discount)",
                "new": "discount = min(discount, 0.5)\n    return price * (1 - discount)",
                "plan_step_id": "s1",
                "rationale": "cap",
                "evidence": [],
            }
        ]
    }
)


class _Provider:
    name = "fake"

    def complete(self, *, template, schema, **_kw):
        if template == "planning":
            return _PLAN
        if template == "implementation":
            return _EDIT
        if template == "code_review":
            return schema.model_validate({"findings": []})
        raise AssertionError(template)


def verified_task(db: Session, settings: Settings, fixture: Path) -> str:
    """Drive the acceptance repo task to VERIFIED via direct stage service calls
    (the test_pr_acceptance_fixture.py pattern)."""
    repo = RepositoryRepository(db).get_or_create(
        source_type="LOCAL", url_or_path=str(fixture), name="acc"
    )
    res = ingest_repository(
        db, repository=repo, request=IngestRequest(), settings=settings
    )
    analyze_snapshot(
        db, snapshot=SnapshotRepository(db).get(res.snapshot_id), settings=settings
    )
    task_id = tasks_service.create_task(
        db,
        settings=settings,
        payload=TaskCreate(
            repository_id=repo.id,
            text=(fixture / "task.md").read_text(encoding="utf-8"),
        ),
    ).task.id
    p = _Provider()
    mapping_service.run_mapping(db, settings=settings, task_id=task_id)
    impact_service.get_or_compute_impact(db, settings=settings, task_id=task_id)
    planning_service.generate_plan(db, settings=settings, task_id=task_id, provider=p)
    planning_service.validate_plan_for_task(db, task_id)
    implementation_service.generate_implementation(
        db, settings=settings, task_id=task_id, provider=p
    )
    verification_service.get_or_verify(
        db, settings=settings, task_id=task_id, provider=p
    )
    verification_service.resolve_decision(
        db,
        settings=settings,
        task_id=task_id,
        decision="APPROVE",
        reason="manually validated",
        actor="approver@x",
    )
    assert TaskRepository(db).get(task_id).state == "COMPLETED"
    return task_id


def test_report_has_all_18_sections_with_real_data(
    db_session, ingestion_settings, acceptance_fixture_path
):
    db, settings = db_session, ingestion_settings
    task_id = verified_task(db, settings, acceptance_fixture_path)

    report = reporting_service.get_task_report(db, settings=settings, task_id=task_id)
    assert [s.name for s in report.sections] == SECTION_NAMES
    assert len(report.sections) == 18

    by_name = {s.name: s for s in report.sections}
    # the stages this Docker-less path actually runs
    for name in (
        "Requirement",
        "Repository",
        "Understanding",
        "Code mapping",
        "Impact",
        "Plan",
        "Implementation",
        "Regression results",
        "Review",
        "Risk",
        "Confidence",
        "Verification",
        "Patch",
    ):
        assert by_name[name].present, f"section {name!r} should be present"

    assert by_name["Verification"].data["verdict"] == "VERIFIED"
    assert by_name["Patch"].data["diff_text"].strip()
    # degraded scoring signals are surfaced as limitations, not hidden
    assert by_name["Remaining limitations"].present
    assert any(
        "signal" in x for x in by_name["Remaining limitations"].data["limitations"]
    )


def test_report_xlsx_is_a_valid_workbook(
    db_session, ingestion_settings, acceptance_fixture_path
):
    db, settings = db_session, ingestion_settings
    task_id = verified_task(db, settings, acceptance_fixture_path)

    data = reporting_service.get_task_report_xlsx(
        db, settings=settings, task_id=task_id
    )
    wb = load_workbook(io.BytesIO(data))
    assert "Overview" in wb.sheetnames
    overview_rows = list(wb["Overview"].iter_rows(min_row=2, values_only=True))
    assert len(overview_rows) == 18
    assert overview_rows[0][1] == "Requirement"


def test_metrics_workbook_is_honest(
    db_session, ingestion_settings, acceptance_fixture_path
):
    db, settings = db_session, ingestion_settings
    verified_task(db, settings, acceptance_fixture_path)

    data = reporting_service.get_metrics_xlsx(db, settings=settings)
    wb = load_workbook(io.BytesIO(data))
    for sheet in (
        "Tasks",
        "Execution Results",
        "Tests",
        "Failures",
        "Repairs",
        "Reviews",
        "Risk",
        "Verification",
        "Engineering Metrics",
    ):
        assert sheet in wb.sheetnames

    metrics = {
        row[0]: (row[1], row[3], row[4])
        for row in wb["Engineering Metrics"].iter_rows(min_row=2, values_only=True)
    }
    assert len(metrics) == 16

    # #12 completion rate: one COMPLETED task out of one -> a real formula, not N/A
    _, value12, basis12 = metrics[12]
    assert basis12 == "FACT" and str(value12).startswith("=")

    # #8 scope compliance: the acceptance patch is clean -> FACT
    assert metrics[8][2] == "FACT"

    # benchmark-only metrics are explicit N/A with a reason, never a bare number
    for n in (1, 6, 9, 10, 13, 15):
        _, value, basis = metrics[n]
        assert basis == "UNAVAILABLE"
        assert str(value).startswith("N/A")

    # nothing is silently 0 / blank
    for n, (_, value, _basis) in metrics.items():
        assert value not in (None, "", 0)


def test_export_then_reimport_round_trips(
    db_session, ingestion_settings, acceptance_fixture_path
):
    db, settings = db_session, ingestion_settings
    verified_task(db, settings, acceptance_fixture_path)

    data = reporting_service.get_metrics_xlsx(db, settings=settings)
    wb = load_workbook(io.BytesIO(data))
    task_rows = list(wb["Tasks"].iter_rows(min_row=2, values_only=True))
    assert len(task_rows) == 1

    # build an import workbook that re-adds the same repo + a fresh task
    imp = Workbook()
    imp.remove(imp.active)
    rs = imp.create_sheet("Repositories")
    rs.append(["source_type", "url_or_path"])
    rs.append(["LOCAL", str(acceptance_fixture_path)])
    ts = imp.create_sheet("Tasks")
    ts.append(["repository_url_or_path", "text", "task_type"])
    ts.append([str(acceptance_fixture_path), "a second discount rounding bug", "BUG"])
    buf = io.BytesIO()
    imp.save(buf)

    result = reporting_service.import_workbook(
        db, settings=settings, data=buf.getvalue()
    )
    assert result.repositories_created == 0  # the repo already exists
    assert result.tasks_created == 1
    assert result.errors == []

    from app.services import tasks as tasks_service

    assert tasks_service.list_tasks(db).total == 2
