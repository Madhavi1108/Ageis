"""Phase 23 integration: POST /reports/import, GET /reports/tasks/{id},
GET /reports/tasks/{id}.xlsx, GET /reports/metrics.xlsx.
"""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.ai.deps import get_ai_provider
from app.ai.provider import MockProvider
from app.core.config import Settings, get_settings
from app.db.session import get_db
from app.main import app
from app.models.base import Base

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_PLAN = {
    "problem_interpretation": "cap the discount at 0.5",
    "assumptions": [],
    "files_to_inspect": ["invoice.py"],
    "files_to_modify": ["invoice.py"],
    "symbols_to_modify": ["invoice.py::calculate_total"],
    "dependencies": [],
    "steps": [
        {
            "id": "s1",
            "description": "clamp",
            "test_intent": "90==50",
            "evidence_refs": [],
        }
    ],
    "test_strategy": {"approach": "boundary"},
    "expected_behavior": "the discount never exceeds 0.5",
    "regression_risks": [],
    "rollback_strategy": "revert invoice.py",
    "source": "AI",
    "confidence": {"value": 0.8, "basis": "INFERENCE"},
    "evidence": [],
}
_EDIT = {
    "edit_ops": [
        {
            "path": "invoice.py",
            "op": "replace",
            "anchor": "return price * (1 - discount)",
            "new": "discount = min(discount, 0.5)\n    return price * (1 - discount)",
            "plan_step_id": "s1",
            "rationale": "cap",
            "evidence": [],
        }
    ]
}


@pytest.fixture
def client(tmp_path, acceptance_fixture_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'reports_test.db'}")
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    provider = MockProvider()

    def _get_db():
        db = SessionLocal()
        try:
            yield db
        finally:
            db.close()

    def _get_settings():
        return Settings(
            ingestion_local_roots=[str(acceptance_fixture_path.parent)],
            artifacts_root=str(tmp_path / "artifacts"),
            _env_file=None,
        )

    app.dependency_overrides[get_db] = _get_db
    app.dependency_overrides[get_settings] = _get_settings
    app.dependency_overrides[get_ai_provider] = lambda: provider
    try:
        with TestClient(app) as c:
            yield c, provider, acceptance_fixture_path
    finally:
        app.dependency_overrides.clear()


def _verified(c: TestClient, provider: MockProvider, fixture) -> str:
    repo_id = c.post(
        "/repositories", json={"source_type": "LOCAL", "url_or_path": str(fixture)}
    ).json()["id"]
    sid = c.post(f"/repositories/{repo_id}/snapshots", json={}).json()["snapshot_id"]
    c.post(f"/repositories/{repo_id}/snapshots/{sid}/analysis", json={})
    task_md = (fixture / "task.md").read_text(encoding="utf-8")
    task_id = c.post("/tasks", json={"repository_id": repo_id, "text": task_md}).json()[
        "task"
    ]["id"]
    c.post("/analysis/map", json={"task_id": task_id})
    c.get(f"/tasks/{task_id}/impact")
    provider.register("planning", task_id, _PLAN)
    c.post(f"/tasks/{task_id}/plan")
    c.post(f"/tasks/{task_id}/plan/validate")
    provider.register("implementation", task_id, _EDIT)
    assert c.post(f"/tasks/{task_id}/changes").status_code == 201
    c.get(f"/tasks/{task_id}/verification")
    r = c.post(
        f"/tasks/{task_id}/verification/decision",
        json={"decision": "APPROVE", "reason": "ok", "actor": "u@x"},
    )
    assert r.status_code == 200 and r.json()["verdict"] == "VERIFIED", r.text
    return task_id


def test_task_report_json_has_18_sections(client):
    c, provider, fixture = client
    task_id = _verified(c, provider, fixture)

    r = c.get(f"/reports/tasks/{task_id}")
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["task_id"] == task_id
    assert [s["number"] for s in body["sections"]] == list(range(1, 19))
    names = {s["name"]: s for s in body["sections"]}
    assert (
        names["Verification"]["present"]
        and names["Verification"]["data"]["verdict"] == "VERIFIED"
    )
    assert names["Tests"]["present"] is False  # Docker-less: not run

    missing = c.get("/reports/tasks/nope")
    assert missing.status_code == 404
    assert missing.json()["code"] == "REPORT_TASK_NOT_FOUND"


def test_task_report_xlsx_downloads(client):
    c, provider, fixture = client
    task_id = _verified(c, provider, fixture)

    r = c.get(f"/reports/tasks/{task_id}.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == _XLSX_MIME
    assert f"task-{task_id}.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb["Overview"].max_row == 19  # header + 18 sections


def test_metrics_xlsx_downloads_with_nine_sheets(client):
    c, provider, fixture = client
    _verified(c, provider, fixture)

    r = c.get("/reports/metrics.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == _XLSX_MIME
    wb = load_workbook(io.BytesIO(r.content))
    for sheet in (
        "Tasks",
        "Execution Results",
        "Tests",
        "Failures",
        "Repairs",
        "Reviews",
        "Risk",
        "Verification",
        "Engineering Metrics",
    ):
        assert sheet in wb.sheetnames
    assert wb["Engineering Metrics"].max_row == 17  # header + 16 metrics


def test_import_creates_tasks_and_reports_bad_rows(client):
    c, provider, fixture = client

    wb = Workbook()
    wb.remove(wb.active)
    rs = wb.create_sheet("Repositories")
    rs.append(["source_type", "url_or_path"])
    rs.append(["LOCAL", str(fixture)])
    ts = wb.create_sheet("Tasks")
    ts.append(["repository_url_or_path", "text", "task_type", "allowed_paths"])
    ts.append(
        [str(fixture), "cap the discount at 0.5 in invoice totals", "BUG", "invoice.py"]
    )
    ts.append([str(fixture), "a different rounding issue in totals", "BUG", ""])
    ts.append(["/no/such/repo", "orphan row", "", ""])  # bad: unknown repo
    buf = io.BytesIO()
    wb.save(buf)

    r = c.post(
        "/reports/import",
        content=buf.getvalue(),
        headers={"content-type": _XLSX_MIME},
    )
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["repositories_created"] == 1
    assert body["tasks_created"] == 2
    assert len(body["errors"]) == 1
    assert body["errors"][0]["sheet"] == "Tasks"
    assert body["errors"][0]["row"] == 4  # header + 2 good rows, then the bad row

    listed = c.get("/tasks").json()
    assert listed["total"] == 2


def test_import_row_uses_the_same_validation_as_the_api(client):
    """A duplicate task row surfaces the API's own TASK_DUPLICATE code."""
    c, provider, fixture = client
    repo_id = c.post(
        "/repositories", json={"source_type": "LOCAL", "url_or_path": str(fixture)}
    ).json()["id"]
    c.post("/tasks", json={"repository_id": repo_id, "text": "cap discount at 0.5"})

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Repositories").append(["source_type", "url_or_path"])
    ts = wb.create_sheet("Tasks")
    ts.append(["repository_url_or_path", "text"])
    ts.append([str(fixture), "cap discount at 0.5"])  # identical -> duplicate
    buf = io.BytesIO()
    wb.save(buf)

    r = c.post(
        "/reports/import", content=buf.getvalue(), headers={"content-type": _XLSX_MIME}
    )
    body = r.json()
    assert body["tasks_created"] == 0
    assert body["errors"][0]["code"] == "TASK_DUPLICATE"
    assert "existing_task_id" in body["errors"][0]["details"]
