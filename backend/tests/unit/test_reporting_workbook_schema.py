"""Unit: the import-workbook parser validates structure before any DB work
(docs/AEGIS_IMPLEMENTATION_PLAN.md Section 31 -- "workbook schema + validation").
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.reporting.errors import WorkbookParseError, WorkbookSchemaError
from app.reporting.excel_import import parse_import_workbook


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _valid_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    rs = wb.create_sheet("Repositories")
    rs.append(["source_type", "url_or_path", "name"])
    rs.append(["LOCAL", "/repos/acc", "acc"])
    ts = wb.create_sheet("Tasks")
    ts.append(["repository_url_or_path", "text", "task_type"])
    ts.append(["/repos/acc", "fix the thing", "BUG"])
    return wb


def test_parses_a_valid_workbook():
    parsed = parse_import_workbook(_bytes(_valid_workbook()))
    assert [r["source_type"] for r in parsed.repositories.rows] == ["LOCAL"]
    assert parsed.repositories.row_numbers == [2]
    assert parsed.tasks.rows[0]["text"] == "fix the thing"


def test_blank_rows_are_skipped():
    wb = _valid_workbook()
    wb["Tasks"].append([None, None, None])
    wb["Tasks"].append(["/repos/acc", "another", "FEATURE"])
    parsed = parse_import_workbook(_bytes(wb))
    assert [r["text"] for r in parsed.tasks.rows] == ["fix the thing", "another"]
    assert parsed.tasks.row_numbers == [2, 4]  # the blank row 3 is skipped


def test_missing_required_sheet_is_rejected():
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Repositories").append(["source_type", "url_or_path"])
    with pytest.raises(WorkbookSchemaError) as exc:
        parse_import_workbook(_bytes(wb))
    assert exc.value.details["missing_sheet"] == "Tasks"


def test_missing_required_column_is_rejected():
    wb = _valid_workbook()
    # rebuild Repositories without url_or_path
    del wb["Repositories"]
    rs = wb.create_sheet("Repositories")
    rs.append(["source_type", "name"])
    rs.append(["LOCAL", "acc"])
    with pytest.raises(WorkbookSchemaError) as exc:
        parse_import_workbook(_bytes(wb))
    assert "url_or_path" in exc.value.details["missing_columns"]


def test_unknown_columns_are_ignored():
    wb = _valid_workbook()
    wb["Tasks"].cell(row=1, column=9, value="unknown_extra")
    wb["Tasks"].cell(row=2, column=9, value="ignored")
    parsed = parse_import_workbook(_bytes(wb))
    assert "unknown_extra" not in parsed.tasks.rows[0]


def test_non_xlsx_bytes_are_rejected():
    with pytest.raises(WorkbookParseError):
        parse_import_workbook(b"not a workbook")
