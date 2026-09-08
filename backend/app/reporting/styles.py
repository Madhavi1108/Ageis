"""Small openpyxl helpers shared by the export builders. Standard (not
write-only) mode -- corpus sizes are small at MVP scale; write-only streaming
for large corpora is a documented Phase 23 open item.
"""

from __future__ import annotations

from typing import Any, Iterable

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")


def write_header_row(ws: Worksheet, columns: Iterable[str]) -> None:
    cols = list(columns)
    ws.append(cols)
    for idx in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"


def append_row(ws: Worksheet, values: Iterable[Any]) -> None:
    ws.append(["" if v is None else v for v in values])


def autofit(ws: Worksheet, *, max_width: int = 80) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            if length > widths.get(cell.column, 0):
                widths[cell.column] = length
    for col_idx, width in widths.items():
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = min(max_width, max(10, width + 2))
