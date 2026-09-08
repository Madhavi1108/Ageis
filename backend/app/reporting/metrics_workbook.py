"""The corpus metrics workbook (GET /reports/metrics.xlsx) -- nine sheets built
from the domain repositories across every task, plus an Engineering Metrics
sheet.

Honesty rule (Spec Section 43): a metric cell is either a real value backed by
DB rows (via a formula referencing the 'Metric Inputs' sheet) or an explicit
'N/A -- <reason>'. Metrics 1/6/9/10/13/15 need the Phase 25 benchmark harness
(gold sets, seeded faults, reference agents, priced-model token accounting) and
render N/A; the rest are computed from persisted rows and are <1.0-honest in a
Docker-less run.
"""

from __future__ import annotations

from io import BytesIO
from statistics import quantiles
from typing import Any

from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.core.config import Settings
from app.reporting import collect
from app.reporting.styles import append_row, autofit, write_header_row
from app.reporting.workbook_schema import EXPORT_SHEETS, METRICS_SHEET
from app.repository.jobs import JobRepository
from app.repository.tasks import TaskRepository

_NON_COUNTING = {"PARTIALLY_SUPPORTED", "INFRA_ERROR"}


def build_metrics_workbook(db: Session, *, settings: Settings) -> bytes:
    tasks = TaskRepository(db).list(limit=100_000)
    bundles = [(t, _bundle(db, settings, t.id)) for t in tasks]

    wb = Workbook()
    wb.remove(wb.active)

    _sheet(wb, "Tasks", (_task_row(t) for t, _ in bundles))
    _sheet(wb, "Execution Results", _execution_rows(bundles))
    _sheet(wb, "Tests", _test_rows(bundles))
    _sheet(wb, "Failures", _failure_rows(bundles))
    _sheet(wb, "Repairs", _repair_rows(bundles))
    _sheet(wb, "Reviews", _review_rows(bundles))
    _sheet(wb, "Risk", _risk_rows(bundles))
    _sheet(wb, "Verification", _verification_rows(bundles))
    _metrics_sheet(wb, db, bundles)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bundle(db: Session, settings: Settings, task_id: str) -> dict[str, Any]:
    return {
        "impl": collect.implementation(db, task_id),
        "tests": collect.tests(db, task_id),
        "executions": collect.executions(db, task_id),
        "failures": collect.failures(db, settings, task_id),
        "repairs": collect.repairs(db, settings, task_id),
        "review": collect.review(db, settings, task_id),
        "scores": collect.scores(db, settings, task_id),
        "verification": collect.verification(db, settings, task_id),
    }


def _sheet(wb: Workbook, name: str, rows: Any) -> None:
    ws = wb.create_sheet(name)
    write_header_row(ws, EXPORT_SHEETS[name])
    for row in rows:
        append_row(ws, row)
    autofit(ws)


# --- raw sheet rows ---------------------------------------------------------


def _task_row(t: Any) -> list[Any]:
    return [
        t.id,
        t.repository_id,
        t.title,
        t.task_type,
        t.priority,
        t.state,
        t.terminal_reason,
        t.created_at.isoformat() if t.created_at else "",
        t.updated_at.isoformat() if t.updated_at else "",
    ]


def _execution_rows(bundles: list[tuple[Any, dict[str, Any]]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for t, b in bundles:
        for ex in b["executions"]:
            results = ex.results or []
            passed = sum(1 for r in results if r.outcome == "PASS")
            failed = sum(1 for r in results if r.outcome in {"FAIL", "ERROR"})
            out.append(
                [
                    t.id,
                    ex.id,
                    ex.version,
                    ex.outcome,
                    ex.exit_code,
                    passed,
                    failed,
                    len(results),
                    ex.duration_ms,
                    ex.reason,
                ]
            )
    return out


def _test_rows(bundles: list[tuple[Any, dict[str, Any]]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for t, b in bundles:
        tg = b["tests"]
        if tg is None:
            continue
        for tc in tg.test_cases:
            out.append(
                [
                    t.id,
                    tg.version,
                    tc.name,
                    tc.path,
                    tc.target_symbol,
                    tc.kind,
                    tc.status,
                    tc.invalid_reason,
                ]
            )
    return out


def _failure_rows(bundles: list[tuple[Any, dict[str, Any]]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for t, b in bundles:
        fa = b["failures"]
        if fa is None:
            continue
        for rec in fa.failures:
            frames = rec.frames or []
            out.append(
                [
                    t.id,
                    fa.execution_id,
                    rec.test_name,
                    rec.failure_type,
                    rec.exception_type,
                    (rec.message or "")[:500],
                    sum(1 for f in frames if f.in_diff),
                ]
            )
    return out


def _repair_rows(bundles: list[tuple[Any, dict[str, Any]]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for t, b in bundles:
        rr = b["repairs"]
        if rr is None:
            continue
        for a in rr.attempts:
            out.append(
                [
                    t.id,
                    rr.outcome,
                    a.iteration,
                    a.outcome,
                    (a.hypothesis or "")[:300],
                    a.failing_before,
                    a.failing_after,
                    a.regression_failures,
                ]
            )
    return out


def _review_rows(bundles: list[tuple[Any, dict[str, Any]]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for t, b in bundles:
        rv = b["review"]
        if rv is None:
            continue
        c = rv.counts_by_severity or {}
        out.append(
            [
                t.id,
                rv.implementation_version,
                bool(rv.blocking),
                c.get("CRITICAL", 0),
                c.get("HIGH", 0),
                c.get("MEDIUM", 0),
                c.get("LOW", 0),
                c.get("INFO", 0),
                ", ".join(rv.static_tools_run or []),
            ]
        )
    return out


def _risk_rows(bundles: list[tuple[Any, dict[str, Any]]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for t, b in bundles:
        if b["scores"] is None:
            continue
        pcs, crs = b["scores"]
        out.append(
            [
                t.id,
                pcs.value,
                pcs.classification,
                pcs.overall_confidence,
                crs.value,
                crs.classification,
                crs.overall_confidence,
                ", ".join(pcs.hard_gate or []),
            ]
        )
    return out


def _verification_rows(bundles: list[tuple[Any, dict[str, Any]]]) -> list[list[Any]]:
    out: list[list[Any]] = []
    for t, b in bundles:
        vr = b["verification"]
        if vr is None:
            continue
        mandatory = [c for c in vr.criteria if c.mandatory]
        mand_pass = sum(1 for c in mandatory if c.verdict == "PASS")
        out.append(
            [
                t.id,
                vr.verdict,
                mand_pass,
                len(mandatory),
                vr.replay_fidelity,
                vr.resulting_state,
                vr.decision.decision if vr.decision else None,
                vr.model_version,
            ]
        )
    return out


# --- Engineering Metrics --------------------------------------------------


def _metrics_sheet(
    wb: Workbook, db: Session, bundles: list[tuple[Any, dict[str, Any]]]
) -> None:
    inputs = _compute_inputs(db, bundles)

    isheet = wb.create_sheet("Metric Inputs")
    write_header_row(isheet, ["num", "numerator", "denominator", "note"])
    row_for: dict[int, int] = {}
    r = 2
    for num in sorted(k for k in inputs if k > 0):
        numerator, denominator, note = inputs[num]
        append_row(isheet, [num, numerator, denominator, note])
        row_for[num] = r
        r += 1
    autofit(isheet)

    ws = wb.create_sheet(METRICS_SHEET)
    write_header_row(ws, EXPORT_SHEETS[METRICS_SHEET])
    latency = inputs.get(-14)
    for spec in _METRIC_SPECS:
        num = spec["num"]
        if num in row_for:
            ir = row_for[num]
            value: Any = (
                f"=IF('Metric Inputs'!C{ir}=0,\"n/a\","
                f"'Metric Inputs'!B{ir}/'Metric Inputs'!C{ir})"
            )
            basis = "FACT"
        elif num == 14 and latency and latency[0] > 0:
            value = latency[2]
            basis = "FACT"
        else:
            value = f"N/A — {spec['na_reason']}"
            basis = "UNAVAILABLE"
        append_row(
            ws, [num, spec["name"], spec["formula"], value, basis, spec["source"]]
        )
    autofit(ws)


def _compute_inputs(
    db: Session, bundles: list[tuple[Any, dict[str, Any]]]
) -> dict[int, tuple[float, float, str]]:
    out: dict[int, tuple[float, float, str]] = {}
    total = len(bundles)

    completed = sum(1 for t, _ in bundles if t.state == "COMPLETED")
    out[12] = (float(completed), float(total), "Task.state == COMPLETED / all tasks")

    with_impl = [(t, b) for t, b in bundles if b["impl"] is not None]
    scope_clean = sum(1 for _, b in with_impl if not b["impl"].scope_violations)
    out[8] = (
        float(scope_clean),
        float(len(with_impl)),
        "latest Implementation has no scope violations / tasks with an implementation",
    )

    accepted = sum(
        1
        for _, b in with_impl
        if b["verification"]
        and b["verification"].verdict == "VERIFIED"
        and not b["impl"].scope_violations
    )
    out[7] = (
        float(accepted),
        float(len(with_impl)),
        "VERIFIED & scope-clean / patches generated",
    )

    tests_total = tests_valid = 0
    for _, b in bundles:
        tg = b["tests"]
        if tg is None:
            continue
        for tc in tg.test_cases:
            tests_total += 1
            if tc.status != "INVALID":
                tests_valid += 1
    out[3] = (
        float(tests_valid),
        float(tests_total),
        "TestCase.status != INVALID / generated",
    )

    passed = executed = 0
    for _, b in bundles:
        for ex in b["executions"]:
            if ex.outcome in _NON_COUNTING:
                continue
            executed += len(ex.results or [])
            passed += sum(1 for r in (ex.results or []) if r.outcome == "PASS")
    out[4] = (
        float(passed),
        float(executed),
        "passed / executed test results (real runs only)",
    )

    alignments: list[float] = []
    for _, b in bundles:
        vr = b["verification"]
        if vr is None:
            continue
        pa = vr.plan_alignment
        steps_total = max(1, pa.steps_total)
        touched = max(1, len(pa.files_touched or []))
        alignments.append(
            (pa.steps_implemented / steps_total)
            * (1 - len(pa.unplanned_files or []) / touched)
        )
    if alignments:
        out[2] = (sum(alignments), float(len(alignments)), "mean per-task alignment")

    entering = repaired = iterations = 0
    for _, b in bundles:
        rr = b["repairs"]
        if rr is None or not rr.attempts:
            continue
        entering += 1
        iterations += len(rr.attempts)
        if rr.outcome == "REPAIRED":
            repaired += 1
    if entering:
        out[5] = (
            float(repaired),
            float(entering),
            "repaired-to-green / tasks entering repair",
        )
        out[11] = (
            float(iterations),
            float(entering),
            "sum(attempts) / tasks entering repair",
        )

    fidelities = [
        b["verification"].replay_fidelity
        for _, b in bundles
        if b["verification"] is not None
        and b["verification"].replay_fidelity is not None
    ]
    if fidelities:
        out[16] = (
            float(sum(fidelities)),
            float(len(fidelities)),
            "mean verification replay_fidelity",
        )

    durations: list[float] = []
    for t, _ in bundles:
        for job in JobRepository(db).list_for_task(t.id):
            if job.type == "RUN_TASK" and job.started_at and job.finished_at:
                durations.append((job.finished_at - job.started_at).total_seconds())
    out[-14] = (float(len(durations)), 0.0, _latency_note(durations))
    return out


def _latency_note(durations: list[float]) -> str:
    if not durations:
        return "no completed RUN_TASK jobs with timestamps"
    if len(durations) == 1:
        return f"P50={durations[0]:.1f}s P95={durations[0]:.1f}s (n=1)"
    qs = quantiles(sorted(durations), n=100, method="inclusive")
    return f"P50={qs[49]:.1f}s P95={qs[94]:.1f}s (n={len(durations)})"


_METRIC_SPECS: list[dict[str, Any]] = [
    {
        "num": 1,
        "name": "Issue->Code Mapping Accuracy",
        "formula": "F1(predicted, gold); recall@k",
        "source": "mapping vs gold file lists",
        "na_reason": "needs gold file lists (Phase 25 benchmark set)",
    },
    {
        "num": 2,
        "name": "Plan->Implementation Alignment",
        "formula": "(steps_impl/steps_total)*(1 - unplanned/max(1,files_touched))",
        "source": "Verification sheet",
        "na_reason": "no verified tasks with a plan-alignment record",
    },
    {
        "num": 3,
        "name": "Test Generation Validity",
        "formula": "tests_valid / tests_generated",
        "source": "Tests sheet",
        "na_reason": "no tests generated",
    },
    {
        "num": 4,
        "name": "Test Pass Rate",
        "formula": "tests_passed / tests_executed",
        "source": "Execution Results sheet",
        "na_reason": "no real (non PARTIALLY_SUPPORTED) executions — needs Docker",
    },
    {
        "num": 5,
        "name": "Autonomous Repair Success Rate",
        "formula": "repaired_to_green / tasks_entering_repair",
        "source": "Repairs sheet",
        "na_reason": "no task entered the repair loop",
    },
    {
        "num": 6,
        "name": "Regression Detection Rate",
        "formula": "injected_regressions_caught / injected_regressions_total",
        "source": "seeded regressions",
        "na_reason": "needs a seeded-regression set (Phase 25)",
    },
    {
        "num": 7,
        "name": "Patch Acceptance Rate",
        "formula": "patches_verified_and_scope_clean / patches_generated",
        "source": "Verification + Risk sheets",
        "na_reason": "no patches generated",
    },
    {
        "num": 8,
        "name": "Scope Compliance",
        "formula": "1 - tasks_with_unjustified_scope_violation / tasks_total",
        "source": "Implementation scope_violations",
        "na_reason": "no tasks with an implementation",
    },
    {
        "num": 9,
        "name": "Review Defect Detection",
        "formula": "seeded_defects_flagged / seeded_defects_total",
        "source": "seeded-defect patches",
        "na_reason": "needs seeded-defect patches (Phase 25)",
    },
    {
        "num": 10,
        "name": "Verification Accuracy / false-complete",
        "formula": "(TP+TN)/total ; false_complete = FP/(FP+TN)",
        "source": "verdicts vs ground truth",
        "na_reason": "needs a labeled verification set (Phase 25)",
    },
    {
        "num": 11,
        "name": "Mean Repair Iterations",
        "formula": "sum(iterations) / tasks_entering_repair",
        "source": "Repairs sheet",
        "na_reason": "no task entered the repair loop",
    },
    {
        "num": 12,
        "name": "Task Completion Rate",
        "formula": "tasks_completed_verified / tasks_submitted",
        "source": "Tasks sheet",
        "na_reason": "no tasks",
    },
    {
        "num": 13,
        "name": "Cost per Verified Task (USD)",
        "formula": "total_priced_model_spend / tasks_verified",
        "source": "token accounting + pricing-table",
        "na_reason": "no priced-model token accounting recorded (Phase 25 pricing-table)",
    },
    {
        "num": 14,
        "name": "Task Latency P50 / P95",
        "formula": "percentiles of wall-clock run->terminal",
        "source": "Job started_at/finished_at",
        "na_reason": "no completed RUN_TASK jobs with timestamps",
    },
    {
        "num": 15,
        "name": "Competitive Resolution-Rate Delta",
        "formula": "aegis_rate - best_reference_agent_rate",
        "source": "AEGIS + reference-agent outcomes",
        "na_reason": "needs reference-agent outcomes (Phase 25)",
    },
    {
        "num": 16,
        "name": "Deterministic Replay Fidelity",
        "formula": "tasks_replaying_to_identical_patch / tasks_sampled",
        "source": "Verification replay_fidelity",
        "na_reason": "no verification rows with a replay-fidelity value",
    },
]
