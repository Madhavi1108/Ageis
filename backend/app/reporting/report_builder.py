"""The 18-section structured task report (Specification Section 33) plus its
xlsx rendering. Pure reads only -- extends app/github/pr_builder.py's approach
(no AI, repository-layer access) to every stage via app/reporting/collect.py.
"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from sqlalchemy.orm import Session

from app.analysis.analyze import build_analysis_result
from app.core.config import Settings
from app.reporting import collect
from app.reporting.errors import ReportTaskNotFoundError
from app.reporting.styles import append_row, autofit, write_header_row
from app.repository.analyses import AnalysisRepository
from app.repository.issues import IssueRepository
from app.repository.pull_requests import PullRequestRepository
from app.repository.repositories import RepositoryRepository
from app.repository.snapshots import SnapshotRepository
from app.repository.tasks import TaskRepository
from app.repository.test_executions import TestExecutionRepository
from app.schemas.report import SECTION_NAMES, ReportSection, TaskReport
from app.schemas.repository import RepositoryRef
from app.services import pr as pr_service
from app.services import tasks as tasks_service


def _dump(obj: Any) -> Any:
    if obj is None:
        return None
    if isinstance(obj, list):
        return [_dump(x) for x in obj]
    if hasattr(obj, "model_dump"):
        return obj.model_dump(mode="json")
    return obj


def _section(
    number: int, present: bool, data: Any = None, note: str = ""
) -> ReportSection:
    name = SECTION_NAMES[number - 1]
    if not present:
        return ReportSection(
            number=number,
            name=name,
            present=False,
            note=note or "not produced for this task",
        )
    return ReportSection(
        number=number, name=name, present=True, note=note, data=_dump(data)
    )


def build_task_report(db: Session, *, settings: Settings, task_id: str) -> TaskReport:
    if TaskRepository(db).get(task_id) is None:
        raise ReportTaskNotFoundError(f"task {task_id} not found")
    task = tasks_service.get_task(db, task_id)  # the API Task schema shape

    repo = RepositoryRepository(db).get(task.repository_id)
    snapshot = (
        SnapshotRepository(db).get(task.snapshot_id) if task.snapshot_id else None
    )
    analysis = (
        AnalysisRepository(db).get_by_snapshot(task.snapshot_id)
        if task.snapshot_id
        else None
    )
    issue_ref = None
    if task.issue_id:
        issue = IssueRepository(db).get(task.issue_id)
        issue_ref = issue.external_ref if issue else None

    impl = collect.implementation(db, task_id)
    scores = collect.scores(db, settings, task_id)
    pr_present = PullRequestRepository(db).get_latest_by_task(task_id) is not None

    sections: list[ReportSection] = [
        _section(
            1,
            True,
            {
                "task_id": task.id,
                "title": task.title,
                "task_type": task.task_type,
                "priority": task.priority,
                "description": task.description,
                "allowed_paths": task.allowed_paths or [],
                "issue_external_ref": issue_ref,
                "created_by": task.created_by,
                "created_at": task.created_at.isoformat(),
            },
        ),
        _section(
            2,
            repo is not None,
            (
                None
                if repo is None
                else {
                    "repository": RepositoryRef.model_validate(
                        repo, from_attributes=True
                    ).model_dump(mode="json"),
                    "snapshot_id": task.snapshot_id,
                    "commit_sha": getattr(snapshot, "commit_sha", None),
                    "branch": getattr(snapshot, "branch", None),
                }
            ),
        ),
        _section(
            3,
            analysis is not None,
            None if analysis is None else build_analysis_result(analysis, job_id=""),
        ),
        _section(4, (m := collect.mapping(db, task_id)) is not None, m),
        _section(5, (i := collect.impact(db, settings, task_id)) is not None, i),
        _section(6, (p := collect.plan(db, task_id)) is not None, p),
        _section(7, impl is not None, impl),
        _section(8, (tg := collect.tests(db, task_id)) is not None, tg),
        _section(9, (fa := collect.failures(db, settings, task_id)) is not None, fa),
        _section(10, (rp := collect.repairs(db, settings, task_id)) is not None, rp),
        _section(11, (rg := collect.regression(db, settings, task_id)) is not None, rg),
        _section(12, (rv := collect.review(db, settings, task_id)) is not None, rv),
        _section(13, scores is not None, None if scores is None else scores[1]),
        _section(14, scores is not None, None if scores is None else scores[0]),
        _section(
            15,
            (vr := collect.verification(db, settings, task_id)) is not None,
            vr,
        ),
        _section(16, impl is not None, None if impl is None else impl.patch),
        _section(
            17,
            pr_present,
            pr_service.get_pr(db, task_id) if pr_present else None,
        ),
    ]
    sections.append(
        _limitations_section(db, task_id, task.state, task.terminal_reason, sections)
    )

    return TaskReport(
        task_id=task.id,
        repository_id=task.repository_id,
        title=task.title,
        outcome=task.state,
        generated_at=datetime.now(timezone.utc),
        sections=sections,
    )


def _limitations_section(
    db: Session,
    task_id: str,
    state: str,
    terminal_reason: str | None,
    prior: list[ReportSection],
) -> ReportSection:
    limits: list[str] = []
    by_name = {s.name: s for s in prior}

    v = by_name.get("Verification")
    if v and v.present and v.data:
        for crit in v.data.get("criteria", []):
            if crit.get("verdict") in {"FAIL", "UNKNOWN"}:
                limits.append(
                    f"verification criterion {crit['name']} = {crit['verdict']}"
                )

    rv = by_name.get("Review")
    if rv and rv.present and rv.data and rv.data.get("blocking"):
        limits.append("code review is blocking (an OPEN CRITICAL/HIGH finding)")

    for score_name in ("Risk", "Confidence"):
        sec = by_name.get(score_name)
        if sec and sec.present and sec.data:
            for sig in sec.data.get("per_signal_contributions", []):
                if sig.get("unavailable_reason"):
                    limits.append(
                        f"{score_name.lower()} signal {sig['name']}: {sig['unavailable_reason']}"
                    )

    for ex in TestExecutionRepository(db).list_for_task(task_id):
        if ex.outcome == "PARTIALLY_SUPPORTED":
            limits.append(
                f"test execution v{ex.version} is PARTIALLY_SUPPORTED "
                f"({ex.reason or 'sandbox unavailable'})"
            )
            break

    rp = by_name.get("Repairs")
    if rp and rp.present and rp.data and rp.data.get("outcome") == "SAFE_STOP":
        ss = rp.data.get("safe_stop") or {}
        limits.append(
            "repair loop stopped safely: " + (ss.get("reason") or "unresolved")
        )

    if state in {"FAILED", "CANCELLED", "PARTIALLY_SUPPORTED"} and terminal_reason:
        limits.append(f"task ended {state}: {terminal_reason}")

    if state != "COMPLETED" and not limits:
        limits.append(f"outcome is {state}; the change is not verified-complete")

    seen: set[str] = set()
    deduped = [x for x in limits if not (x in seen or seen.add(x))]
    return ReportSection(
        number=18,
        name=SECTION_NAMES[17],
        present=bool(deduped),
        note="" if deduped else "no limitations recorded",
        data={"limitations": deduped},
    )


def render_task_report_workbook(report: TaskReport) -> bytes:
    wb = Workbook()
    overview = wb.active
    overview.title = "Overview"
    write_header_row(overview, ["#", "Section", "Status", "Summary"])
    for sec in report.sections:
        append_row(
            overview,
            [
                sec.number,
                sec.name,
                "present" if sec.present else "absent",
                _summarize(sec),
            ],
        )
    autofit(overview)

    meta = wb.create_sheet("Report")
    write_header_row(meta, ["field", "value"])
    for k, val in (
        ("task_id", report.task_id),
        ("repository_id", report.repository_id),
        ("title", report.title),
        ("outcome", report.outcome),
        ("generated_at", report.generated_at.isoformat()),
    ):
        append_row(meta, [k, val])
    autofit(meta)

    for sec in report.sections:
        if not sec.present or sec.data is None:
            continue
        ws = wb.create_sheet(_sheet_title(sec.number, sec.name))
        write_header_row(ws, ["key", "value"])
        for key, value in _flatten(sec.data):
            append_row(ws, [key, value])
        autofit(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_title(number: int, name: str) -> str:
    raw = f"{number:02d} {name}"
    for ch in "[]:*?/\\":
        raw = raw.replace(ch, " ")
    return raw[:31]


def _summarize(sec: ReportSection) -> str:
    if not sec.present:
        return sec.note or "absent"
    data = sec.data
    if isinstance(data, dict):
        for key in ("verdict", "outcome", "classification", "overall_confidence"):
            if key in data:
                return f"{key}={data[key]}"
        if "limitations" in data:
            return f"{len(data['limitations'])} limitation(s)"
        return f"{len(data)} field(s)"
    if isinstance(data, list):
        return f"{len(data)} item(s)"
    return str(data)[:120]


def _flatten(value: Any, prefix: str = "") -> list[tuple[str, Any]]:
    out: list[tuple[str, Any]] = []
    if isinstance(value, dict):
        for k, v in value.items():
            out.extend(_flatten(v, f"{prefix}{k}."))
    elif isinstance(value, list):
        if not value:
            out.append((prefix.rstrip("."), "[]"))
        for i, v in enumerate(value):
            out.extend(_flatten(v, f"{prefix}{i}."))
    else:
        out.append((prefix.rstrip("."), "" if value is None else value))
    return out
